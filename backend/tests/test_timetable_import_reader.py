"""The reader over small in-memory workbooks shaped like real faculty sheets."""
from datetime import date, time

import openpyxl

from modules.timetable_import.reader import read_sheet


def _sheet(rows, merges=()):
    ws = openpyxl.Workbook().active
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            if value is not None:
                ws.cell(r, c, value)
    for rng in merges:
        ws.merge_cells(rng)
    return ws


def _summary(result):
    return [(e.weekday, e.start, e.end, e.codes, e.venues) for e in result.entries]


def test_lecture_grid_with_day_blocks():
    ws = _sheet([
        ["FACULTY OF LAW"],
        ["MONDAY"],
        ["Venue", "Faculty", "8:00am - 9:30am", "9:30am - 11:00am"],
        ["LLM RM", "Law", "PUL 513", "PUL 513"],
        ["TUESDAY"],
        ["Venue", "Faculty", "8:00am - 9:30am", "9:30am - 11:00am"],
        ["LAW ANNEX", "Law", None, "JIL 412/402"],
    ])
    result = read_sheet(ws)
    assert result.shape == "grid"
    assert _summary(result) == [
        ("Monday", time(8), time(11), ["PUL513"], ["LLM RM", "Law"]),
        ("Tuesday", time(9, 30), time(11), ["JIL412", "JIL402"], ["LAW ANNEX", "Law"]),
    ]


def test_exam_grid_with_days_across_the_top():
    ws = _sheet([
        ["WEEK 1", "DAY 1: MONDAY, 10 AUGUST 2026", None, "DAY 2: TUESDAY, 11 AUGUST 2026"],
        ["VENUE", "9am - 12pm", "12pm - 3pm", "9am - 12pm"],
        ["A/TH", "PHL 105", "HSS 325 (250)", "ENG 320"],
    ])
    entries = read_sheet(ws).entries
    assert [(e.weekday, e.date, e.week, e.start, e.end, e.codes) for e in entries] == [
        ("Monday", date(2026, 8, 10), 1, time(9), time(12), ["PHL105"]),
        ("Monday", date(2026, 8, 10), 1, time(12), time(15), ["HSS325"]),
        ("Tuesday", date(2026, 8, 11), 1, time(9), time(12), ["ENG320"]),
    ]
    assert entries[1].notes == ["250"]


def test_day_in_first_cell_with_room_in_brackets():
    ws = _sheet([
        ["DAYS/TIME", "8:00am-10:00am", "10:00am-12:00pm"],
        ["MONDAY", "CMS 102 (MCB 5)", "MAS 329 (MCB 3)"],
    ])
    entries = read_sheet(ws).entries
    assert [(e.weekday, e.start, e.codes, e.notes, e.venues) for e in entries] == [
        ("Monday", time(8), ["CMS102"], ["MCB 5"], []),
        ("Monday", time(10), ["MAS329"], ["MCB 3"], []),
    ]


def test_merged_time_header_is_shared_between_its_columns():
    ws = _sheet([
        ["VENUE", "MONDAY"],
        [None, "8 - 10am", None, "10 - 12pm", None],
        ["G05", "GEO101", None, "GEO 291", "GEO 201"],
    ], merges=["B2:C2", "D2:E2", "B3:C3"])
    assert [(e.start, e.end, e.codes) for e in read_sheet(ws).entries] == [
        (time(8), time(10), ["GEO101"]),
        (time(10), time(11), ["GEO291"]),
        (time(11), time(12), ["GEO201"]),
    ]


def test_list_sheet():
    ws = _sheet([
        ["Lecture list"],
        [],
        ["Day", "Start", "End", "Venue", "Prefix", "Course", "Level"],
        ["Monday", time(8), time(10), "G407", "LAG", "MET133", 100],
        ["Tue", time(14), time(16), "Hall 1A", None, "SOC 201", 200],
        ["Note: rooms may change"],
    ])
    result = read_sheet(ws)
    assert result.shape == "list"
    assert [(e.weekday, e.start, e.end, e.codes, e.venues, e.level) for e in result.entries] == [
        ("Monday", time(8), time(10), ["LAG-MET133"], ["G407"], "100"),
        ("Tuesday", time(14), time(16), ["SOC201"], ["Hall 1A"], "200"),
    ]
    assert result.warnings == ["1 row(s) were missing a day, time or course"]


def test_sheet_without_a_timetable():
    ws = _sheet([["Year", "RGDP"], [1990, 21680.2]])
    result = read_sheet(ws)
    assert (result.shape, result.entries) == ("none", [])
